# -*- coding: utf-8 -*-
"""论文修改核验：对比两个 docx，出具"只改了该改的"证据报告。

用法：
    python verify_docx.py 旧.docx 新.docx

检查项：
    1. 结构计数前后对比（段落/表格/图片/公式/引文）
    2. 关键数字多重集合前后比对（差异逐条列出）
    3. 逐段 diff（差异段落索引+预览）
    4. 占位符/乱码/LaTeX 残片扫描
    5. 交叉引用完整性（见图/见表/式(x)/见x.x节 的目标是否存在）
    6. 引文编号连续性（1..N）

只读，不修改任何文件。依赖：python-docx。
"""
import io
import re
import sys
import difflib
from collections import Counter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from docx import Document
from docx.oxml.ns import qn

CITE_RE = re.compile(r"\[(\d+(?:[,\-\u2013]\s*\d+)*)\]")
NUM_RE = re.compile(r"\d+\.\d+|\d+")
PLACEHOLDER_RES = [
    ("REPLACE_ME", re.compile(r"REPLACE_ME")),
    ("TODO/FIXME", re.compile(r"TODO|FIXME")),
    ("待确认占位", re.compile(r"\u3010\u5f85[^\u3011]*\u3011")),
    ("LaTeX残片", re.compile(r"\\ln|\\frac|\\%|_\{|\^\{")),
    ("乱码", re.compile(r"\ufffd")),
]
CAPTION_RE = re.compile(r"^\s*(\u56fe|\u8868)\s*([0-9A-Z]+-[0-9]+)")
REF_FIG_RE = re.compile(r"(?:\u89c1|\u5982)\s*(\u56fe|\u8868)\s*([0-9A-Z]+-[0-9]+)")
REF_EQ_RE = re.compile(r"\u5f0f\s*[((\uff08)]([0-9A-Z]+-[0-9]+)[)\uff09]")
REF_SEC_RE = re.compile(r"\u89c1\s*(\d+(?:\.\d+)+)\s*\u8282")


def para_text(p):
    """段落全文（含 hyperlink 内的 run）。"""
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def body_paragraph_texts(doc):
    """正文顶层段落文本（不含表格单元格），按 body 顺序。"""
    return [para_text(p) for p in doc.paragraphs]


def counts(doc):
    body = doc.element.body
    n_para = len(body.findall(qn("w:p")))
    n_tbl = len(body.findall(qn("w:tbl")))
    n_img = len(list(body.iter(qn("a:blip"))))
    n_eq = len(list(body.iter(qn("m:oMath"))))
    return n_para, n_tbl, n_img, n_eq


def citation_numbers(texts):
    nums = []
    for t in texts:
        for m in CITE_RE.finditer(t):
            inner = m.group(1)
            if not re.fullmatch(r"\d+(?:[,\-\u2013]\s*\d+)*", inner):
                continue
            for part in re.split(r"[,\u2013]", inner):
                part = part.strip().rstrip("-").strip()
                if part.isdigit():
                    nums.append(int(part))
    return nums


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    old_doc, new_doc = Document(sys.argv[1]), Document(sys.argv[2])
    old_texts = body_paragraph_texts(old_doc)
    new_texts = body_paragraph_texts(new_doc)
    print("=" * 62)
    print("核验对象：%s  ->  %s" % (sys.argv[1], sys.argv[2]))
    print("=" * 62)

    # 1. 结构计数
    print("\n[1] 结构计数前后对比")
    print("%-12s %8s %8s %s" % ("项", "旧", "新", "判定"))
    for name, o, n in zip(("段落", "表格", "图片", "公式"), counts(old_doc), counts(new_doc)):
        print("%-12s %8d %8d %s" % (name, o, n, "一致" if o == n else "★不一致"))
    co, cn = citation_numbers(old_texts), citation_numbers(new_texts)
    print("%-12s %8d %8d %s" % ("引文标注", len(co), len(cn),
                                 "一致" if len(co) == len(cn) else "★不一致"))

    # 2. 关键数字多重集合
    print("\n[2] 关键数字多重集合比对")
    old_nums = Counter(NUM_RE.findall(" ".join(old_texts)))
    new_nums = Counter(NUM_RE.findall(" ".join(new_texts)))
    only_old = old_nums - new_nums
    only_new = new_nums - old_nums
    if not only_old and not only_new:
        print("数字集合完全一致")
    else:
        print("仅旧版有：%s" % (", ".join("%s×%d" % (k, v) for k, v in sorted(only_old.items())[:40]) or "无"))
        print("仅新版有：%s" % (", ".join("%s×%d" % (k, v) for k, v in sorted(only_new.items())[:40]) or "无"))
        print("（以上差异须逐条解释；语言/格式轮次应为空）")

    # 3. 逐段 diff
    print("\n[3] 逐段 diff（去首尾空白后比较）")
    norm_old = [t.strip() for t in old_texts]
    norm_new = [t.strip() for t in new_texts]
    sm = difflib.SequenceMatcher(None, norm_old, norm_new, autojunk=False)
    n_diff = 0
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        n_diff += (i2 - i1) + (j2 - j1)
        print("  %s 旧[%d:%d] 新[%d:%d]" % (tag, i1, i2, j1, j2))
        for k in range(i1, min(i2, i1 + 3)):
            print("    旧段%d: %s" % (k, norm_old[k][:60]))
        for k in range(j1, min(j2, j1 + 3)):
            print("    新段%d: %s" % (k, norm_new[k][:60]))
    print("  差异段数：%d（须与计划改动位置一致）" % n_diff)

    # 4. 占位符/残片
    print("\n[4] 占位符/乱码/LaTeX 残片扫描")
    hits = 0
    for label, rx in PLACEHOLDER_RES:
        found = [i for i, t in enumerate(new_texts) if rx.search(t)]
        if found:
            hits += len(found)
            print("  ★%s：%d 处，段号 %s" % (label, len(found), found[:20]))
    if not hits:
        print("  未发现残留")

    # 5. 交叉引用
    print("\n[5] 交叉引用完整性（新版）")
    figs, tbls, secs, eqs = set(), set(), set(), set()
    for p in new_doc.paragraphs:
        t = para_text(p)
        style = (p.style.name or "") if p.style is not None else ""
        m = CAPTION_RE.match(t)
        if m:
            (figs if m.group(1) == "\u56fe" else tbls).add(m.group(2))
            continue
        if style.lower().startswith("heading"):
            m = re.match(r"^\s*(\d+(?:\.\d+)*)", t)
            if m:
                secs.add(m.group(1))
        # 公式编号：制表位版式下编号在段尾，按"含 oMath 的段落"提取 (x-y)（兼容全角括号）
        if p._p.findall(".//" + qn("m:oMath")):
            for m in re.finditer(r"[(\uff08]([0-9A-Z]+-[0-9]+)[)\uff09]", t):
                eqs.add(m.group(1))
    full_new = "\n".join(new_texts)
    missing = []
    for m in REF_FIG_RE.finditer(full_new):
        target = m.group(2)
        pool = figs if m.group(1) == "\u56fe" else tbls
        if target not in pool:
            missing.append("%s%s" % (m.group(1), target))
    for m in REF_EQ_RE.finditer(full_new):
        if m.group(1) not in eqs:
            missing.append("式(%s)" % m.group(1))
    for m in REF_SEC_RE.finditer(full_new):
        if m.group(1) not in secs:
            missing.append("%s节" % m.group(1))
    if missing:
        print("  ★悬空引用 %d 处：%s" % (len(missing), ", ".join(sorted(set(missing))[:30])))
    else:
        print("  无悬空引用（图 %d / 表 %d / 式 %d / 节 %d 个目标）"
              % (len(figs), len(tbls), len(eqs), len(secs)))

    # 6. 引文编号连续性
    print("\n[6] 引文编号连续性（新版）")
    if cn:
        mx = max(cn)
        gaps = [n for n in range(1, mx + 1) if n not in set(cn)]
        print("  引文编号范围 1..%d，共 %d 处；缺失编号：%s"
              % (mx, len(cn), gaps if gaps else "无"))
        if gaps:
            print("  （缺失编号可能来自区间展开遗漏或条目被删，须逐条确认）")
    else:
        print("  未检测到引文标注")

    print("\n" + "=" * 62)
    print("核验完毕。语言/格式轮次的预期：[1]全一致、[2]空差异、"
          "[3]仅计划位置、[4]无残留、[5]无悬空、[6]无缺失。")
    print("=" * 62)


if __name__ == "__main__":
    main()
